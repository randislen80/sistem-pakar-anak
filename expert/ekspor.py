"""
Ekspor statistik perkembangan anak ke Excel (openpyxl).

Menerima struktur data biasa dari modul statistik, sehingga dapat diuji
tanpa basis data.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .inference import KESIMPULAN_INFO, STATUS_INFO, STATUS_URUT
from .statistik import peta_hasil, rata_rata_rekaman, tanggal_indonesia

BIRU = "2A78D6"
ISI_KEPALA = PatternFill("solid", fgColor=BIRU)
FONT_KEPALA = Font(bold=True, color="FFFFFF")
FONT_JUDUL = Font(bold=True, size=14)
FONT_SUB = Font(italic=True, color="6B7280")
GARIS = Side(style="thin", color="D9DEE5")
BINGKAI = Border(left=GARIS, right=GARIS, top=GARIS, bottom=GARIS)
ISI_STATUS = {
    "sesuai": PatternFill("solid", fgColor="D8F3D8"),
    "perlu perhatian": PatternFill("solid", fgColor="FEF0CC"),
    "terlambat": PatternFill("solid", fgColor="FBE0D4"),
    "belum terlihat": PatternFill("solid", fgColor="F6D5D5"),
}


def _judul(ws, judul, keterangan, sekolah):
    ws["A1"] = judul
    ws["A1"].font = FONT_JUDUL
    ws["A2"] = sekolah.get("nama", "")
    ws["A3"] = keterangan
    ws["A3"].font = FONT_SUB
    ws["A4"] = f"Diekspor: {tanggal_indonesia(date.today(), pendek=False)}"
    ws["A4"].font = FONT_SUB
    return 6


def _kepala(ws, baris, kolom_list):
    for i, teks in enumerate(kolom_list, start=1):
        c = ws.cell(row=baris, column=i, value=teks)
        c.fill, c.font, c.border = ISI_KEPALA, FONT_KEPALA, BINGKAI
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[baris].height = 30


def _isi(ws, baris, nilai_list, persen_kolom=()):
    for i, v in enumerate(nilai_list, start=1):
        c = ws.cell(row=baris, column=i, value=v)
        c.border = BINGKAI
        if i in persen_kolom and isinstance(v, (int, float)):
            c.number_format = "0.0"
            c.alignment = Alignment(horizontal="center")
    return baris + 1


def _lebar(ws, lebar):
    for i, w in enumerate(lebar, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def excel_statistik(ringkasan: dict, rekaman: list[dict], domains: list[str], keterangan: str,
                    sekolah: dict) -> bytes:
    wb = Workbook()

    # ------------------------------------------------------------------ Ringkasan
    ws = wb.active
    ws.title = "Ringkasan"
    b = _judul(ws, "Statistik Perkembangan Anak", keterangan, sekolah)
    ws.cell(row=b, column=1, value="Indikator").font = Font(bold=True)
    b += 1
    kpi = [
        ("Jumlah anak yang didiagnosis", ringkasan["jumlah_anak"]),
        ("Jumlah diagnosis", ringkasan["jumlah_diagnosis"]),
        ("Rata-rata capaian terakhir (%)", ringkasan["rata_rata"]),
        ("Anak berkembang sesuai harapan", ringkasan["jumlah_sesuai_harapan"]),
        ("Persentase sesuai harapan (%)", ringkasan["persen_sesuai_harapan"]),
        ("Anak yang perlu perhatian", ringkasan["jumlah_perlu_perhatian"]),
    ]
    for label, nilai in kpi:
        b = _isi(ws, b, [label, nilai])
    b += 1

    ws.cell(row=b, column=1, value="Capaian per domain (diagnosis terakhir tiap anak)").font = Font(bold=True)
    b += 1
    awal_tabel = b
    _kepala(ws, b, ["Domain", "Rata-rata capaian (%)", "Jumlah anak"]
            + [STATUS_INFO[s]["label"] for s in STATUS_URUT])
    b += 1
    sebaran = {x["domain"]: x for x in ringkasan["sebaran_status"]}
    for x in ringkasan["rata_domain"]:
        sb = sebaran.get(x["domain"], {"jumlah": {}})
        b = _isi(ws, b, [x["domain"], x["rata"], x["n"]] + [sb["jumlah"].get(s, 0) for s in STATUS_URUT],
                 persen_kolom=(2,))
    akhir_tabel = b - 1

    if ringkasan["rata_domain"]:
        grafik = BarChart()
        grafik.type = "bar"
        grafik.title = "Rata-rata capaian per domain (%)"
        grafik.y_axis.scaling.min = 0
        grafik.y_axis.scaling.max = 100
        grafik.legend = None
        data = Reference(ws, min_col=2, min_row=awal_tabel, max_row=akhir_tabel)
        kategori = Reference(ws, min_col=1, min_row=awal_tabel + 1, max_row=akhir_tabel)
        grafik.add_data(data, titles_from_data=True)
        grafik.set_categories(kategori)
        grafik.height, grafik.width = 7.5, 16
        ws.add_chart(grafik, f"J{awal_tabel}")

    b += 1
    ws.cell(row=b, column=1, value="Sebaran kesimpulan").font = Font(bold=True)
    b += 1
    _kepala(ws, b, ["Kesimpulan", "Jumlah anak", "Persentase (%)"])
    b += 1
    for k in ringkasan["sebaran_kesimpulan"]:
        b = _isi(ws, b, [k["label"], k["jumlah"], k["persen"]], persen_kolom=(3,))
    _lebar(ws, [44, 20, 14, 14, 16, 14, 16])

    # ------------------------------------------------------------ Data diagnosis
    ws = wb.create_sheet("Data Diagnosis")
    b = _judul(ws, "Data Seluruh Diagnosis", keterangan, sekolah)
    kolom = ["No", "Tanggal", "Nama anak", "Kelompok", "Jenis kelamin", "Usia"] + domains + ["Rata-rata (%)", "Kesimpulan"]
    _kepala(ws, b, kolom)
    ws.freeze_panes = ws.cell(row=b + 1, column=4)
    awal = b
    b += 1
    for no, r in enumerate(sorted(rekaman, key=lambda x: (x["tanggal"], x.get("id") or 0)), start=1):
        peta = peta_hasil(r["hasil"])
        kel = r.get("kelompok") or ""
        jk = {"L": "Laki-laki", "P": "Perempuan"}.get(r.get("jenis_kelamin") or "", "")
        nilai = [no, r["tanggal"], r["nama"], f"Kelompok {kel}" if kel else "", jk, r.get("usia") or ""]
        nilai += [peta[d].get("persen") if d in peta else None for d in domains]
        kunci = r.get("kesimpulan") or ""
        nilai += [rata_rata_rekaman(r), KESIMPULAN_INFO.get(kunci, {}).get("label", kunci)]
        persen_kolom = tuple(range(7, 7 + len(domains) + 1))
        _isi(ws, b, nilai, persen_kolom=persen_kolom)
        ws.cell(row=b, column=2).number_format = "DD-MM-YYYY"
        for j, d in enumerate(domains):
            isi = ISI_STATUS.get(peta.get(d, {}).get("status"))
            if isi:
                ws.cell(row=b, column=7 + j).fill = isi
        b += 1
    if b - 1 > awal:
        ws.auto_filter.ref = f"A{awal}:{get_column_letter(len(kolom))}{b - 1}"
    _lebar(ws, [5, 12, 26, 10, 12, 16] + [13] * len(domains) + [13, 26])
    b += 1
    ws.cell(row=b, column=2, value="Warna sel domain:").font = Font(bold=True)
    for i, s in enumerate(STATUS_URUT):
        c = ws.cell(row=b, column=3 + i, value=STATUS_INFO[s]["label"])
        c.fill = ISI_STATUS[s]

    # ----------------------------------------------------------- Perlu perhatian
    ws = wb.create_sheet("Perlu Perhatian")
    b = _judul(ws, "Anak yang Perlu Perhatian", keterangan, sekolah)
    _kepala(ws, b, ["No", "Nama anak", "Kelompok", "Diagnosis terakhir", "Rata-rata (%)",
                    "Domain Terlambat / Belum Terlihat", "Kesimpulan"])
    b += 1
    for no, a in enumerate(ringkasan["perlu_perhatian"], start=1):
        domain_teks = ", ".join(f"{d['domain']} ({d['label']})" for d in a["domain_bermasalah"])
        b = _isi(ws, b, [no, a["nama"], f"Kelompok {a['kelompok']}" if a["kelompok"] else "", a["tanggal"], a["rata_rata"], domain_teks,
                         KESIMPULAN_INFO.get(a["kesimpulan"], {}).get("label", "")], persen_kolom=(5,))
        ws.cell(row=b - 1, column=4).number_format = "DD-MM-YYYY"
        ws.cell(row=b - 1, column=6).alignment = Alignment(wrap_text=True, vertical="top")
    _lebar(ws, [5, 26, 10, 16, 13, 60, 26])

    # -------------------------------------------------------------- Tren bulanan
    ws = wb.create_sheet("Tren Bulanan")
    b = _judul(ws, "Tren Rata-rata Capaian per Bulan", keterangan, sekolah)
    tren = ringkasan["tren"]
    _kepala(ws, b, ["Bulan", "Jumlah diagnosis", "Rata-rata keseluruhan (%)"] + domains)
    b += 1
    for i, label in enumerate(tren["label"]):
        b = _isi(ws, b, [label, tren["jumlah"][i], tren["keseluruhan"][i]]
                 + [tren["per_domain"][d][i] for d in domains],
                 persen_kolom=tuple(range(3, 4 + len(domains))))
    _lebar(ws, [12, 12, 16] + [13] * len(domains))

    # ------------------------------------------------------------------ Perubahan
    ws = wb.create_sheet("Perubahan")
    b = _judul(ws, "Perubahan Capaian (diagnosis pertama ke terakhir)", keterangan, sekolah)
    per = ringkasan["perubahan"]
    for label, nilai in [("Anak dengan 2 diagnosis atau lebih", per["jumlah_anak"]),
                         ("Capaian naik", per["naik"]), ("Capaian tetap", per["tetap"]),
                         ("Capaian turun", per["turun"]), ("Rata-rata perubahan (poin)", per["rata_selisih"])]:
        b = _isi(ws, b, [label, nilai])
    b += 1
    _kepala(ws, b, ["Nama anak", "Jumlah diagnosis", "Capaian awal (%)", "Capaian akhir (%)", "Perubahan (poin)"])
    b += 1
    for d in per["detail"]:
        b = _isi(ws, b, [d["nama"], d["jumlah_diagnosis"], d["awal"], d["akhir"], d["selisih"]],
                 persen_kolom=(3, 4, 5))
    b += 1
    _kepala(ws, b, ["Domain", "Rata-rata perubahan (poin)", "Jumlah anak"])
    b += 1
    for d in per["per_domain"]:
        b = _isi(ws, b, [d["domain"], d["selisih"], d["n"]], persen_kolom=(2,))
    _lebar(ws, [36, 18, 16, 16, 16])

    keluaran = BytesIO()
    wb.save(keluaran)
    return keluaran.getvalue()
